#!/usr/bin/env python3
"""Download the current Reference Medicine inventory + order-form XLSX.

Source page: https://www.referencemedicine.com/inventory-all-specimens
The page contains one "Download inventory" anchor pointing to a Webflow-CDN
hosted XLSX (~4 MB, ~14,625 specimens x 40 columns) with real Tier+Fee pricing.

URL-rotation gotcha
-------------------
The XLSX is uploaded to Webflow's asset CDN at a hashed path of the form:

  https://cdn.prod.website-files.com/<site_id>/<asset_hash>_<filename>.xlsx

Reference Medicine refreshes the file roughly every two weeks. Each refresh
gets a NEW <asset_hash> and the filename embeds the build date (e.g.
"18-MAY-2026 Reference Medicine inventory & order form.xlsx"), so the FULL
URL rotates. **Never hardcode the URL** — always re-resolve it from the page
HTML right before downloading. This script does exactly that.

The page itself is server-rendered (Webflow + Cloudflare), so the href is
visible in raw HTML — no JS execution, no auth, no anti-bot challenge as of
this writing.

XLSX structure (sheet "All specimens")
--------------------------------------
- Rows 1-9: preamble / instructions ("To place your order: ...").
- Row 9: section banner ("Specimen details" / "Donor details").
- Row 10: column headers (40 columns).
- Rows 11..end: one specimen per row (~14,625 rows).
- Column A ("Mark requested") is the buyer checkbox column — empty by
  default; buyers tick it, save the file, and email it back.
- Column 7 = Tier (1-5), Column 8 = Fee (USD).

Usage
-----
    python scripts/fetch_reference_medicine_inventory.py

Dependencies: requests, openpyxl (stdlib otherwise).
"""

from __future__ import annotations

import re
import sys
from datetime import date
from html.parser import HTMLParser
from pathlib import Path

import requests
from openpyxl import load_workbook

PAGE_URL = "https://www.referencemedicine.com/inventory-all-specimens"
USER_AGENT = (
    "Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7) "
    "AppleWebKit/537.36 (KHTML, like Gecko) Chrome/120.0.0.0 Safari/537.36"
)

# Repo layout: scripts/<this>.py  ->  ../server/biobanks/reference_medicine/inventory/
REPO_ROOT = Path(__file__).resolve().parent.parent
DEST_DIR = REPO_ROOT / "server" / "biobanks" / "reference_medicine" / "inventory"


class _XlsxHrefFinder(HTMLParser):
    """Collect every href ending in .xlsx (case-insensitive, ignoring query)."""

    def __init__(self) -> None:
        super().__init__()
        self.hrefs: list[str] = []

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        if tag.lower() != "a":
            return
        for k, v in attrs:
            if k.lower() == "href" and v and re.search(r"\.xlsx(\?|$)", v, re.IGNORECASE):
                self.hrefs.append(v)


def resolve_xlsx_url(page_url: str = PAGE_URL) -> str:
    """Fetch the inventory page and return the current XLSX download URL.

    Raises a clean error if the link can't be found (so callers can fail loud).
    """
    resp = requests.get(page_url, headers={"User-Agent": USER_AGENT}, timeout=30)
    resp.raise_for_status()

    finder = _XlsxHrefFinder()
    finder.feed(resp.text)

    # Fallback: regex over raw HTML in case an .xlsx URL lives outside an <a>.
    candidates = list(finder.hrefs)
    candidates += re.findall(r'https?://[^\s"\'<>]+\.xlsx[^\s"\'<>]*', resp.text)

    # Dedupe, preserve order.
    seen: set[str] = set()
    unique = [u for u in candidates if not (u in seen or seen.add(u))]

    if not unique:
        raise RuntimeError(
            f"No .xlsx link found on {page_url}. The page layout may have "
            "changed, or Reference Medicine may have moved the download "
            "behind a form. Inspect the page manually before patching."
        )

    # Prefer the Webflow CDN ("website-files.com") if multiple are present.
    cdn = [u for u in unique if "website-files.com" in u]
    return cdn[0] if cdn else unique[0]


def download(url: str, dest: Path) -> int:
    """Stream-download `url` to `dest`, return byte count."""
    dest.parent.mkdir(parents=True, exist_ok=True)
    with requests.get(url, headers={"User-Agent": USER_AGENT}, timeout=120, stream=True) as r:
        r.raise_for_status()
        size = 0
        with dest.open("wb") as f:
            for chunk in r.iter_content(chunk_size=64 * 1024):
                if chunk:
                    f.write(chunk)
                    size += len(chunk)
    return size


def summarize(xlsx_path: Path) -> tuple[int, str]:
    """Return (specimen_row_count, schema_summary) for the 'All specimens' sheet."""
    wb = load_workbook(xlsx_path, read_only=True, data_only=True)
    if "All specimens" not in wb.sheetnames:
        raise RuntimeError(
            f"Expected 'All specimens' sheet in {xlsx_path.name}; "
            f"got {wb.sheetnames}. Schema may have changed."
        )
    ws = wb["All specimens"]

    # Header is on row 10 (rows 1-9 are preamble/banner).
    header_row = next(ws.iter_rows(min_row=10, max_row=10, values_only=True))
    headers = [
        (str(h).replace("\n", " ").strip() if h is not None else f"col{i}")
        for i, h in enumerate(header_row)
    ]

    # Data rows: anything from row 11 onward with a non-empty RM case ID (col 3).
    count = sum(1 for row in ws.iter_rows(min_row=11, values_only=True) if row[3])

    # One-line schema summary: a few load-bearing columns for the voice agent.
    key_cols = [headers[i] for i in (0, 3, 6, 7, 8, 13, 17, 30, 31)]
    schema_summary = (
        f"{len(headers)} cols; key fields = "
        + ", ".join(key_cols)
        + " (Tier 1-5; Fee in USD; col A is buyer checkbox)"
    )
    return count, schema_summary


def main() -> int:
    try:
        url = resolve_xlsx_url()
    except Exception as e:
        print(f"ERROR: could not resolve XLSX URL: {e}", file=sys.stderr)
        return 1

    dest = DEST_DIR / f"{date.today().isoformat()}.xlsx"
    try:
        nbytes = download(url, dest)
    except Exception as e:
        print(f"ERROR: download failed: {e}", file=sys.stderr)
        return 2

    try:
        rows, schema = summarize(dest)
    except Exception as e:
        print(f"ERROR: XLSX inspection failed: {e}", file=sys.stderr)
        return 3

    print(f"source_url: {url}")
    print(f"dest:       {dest}")
    print(f"size:       {nbytes:,} bytes ({nbytes / 1_048_576:.2f} MiB)")
    print(f"rows:       {rows:,} specimens")
    print(f"schema:     {schema}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
